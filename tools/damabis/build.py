# -*- coding: utf-8 -*-
import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.utils import get_column_letter as C
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, DataBarRule, ColorScaleRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.drawing.image import Image as XLImage
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.comments import Comment

BASE="/tmp/claude-0/-home-user-pagina-web/38b1380c-584b-55c2-8699-038270dfa567/scratchpad/damabis"
OUT=BASE+"/DAMABIS_Control_Inventario_y_Ventas.xlsx"
LOGO=BASE+"/ajolote.png"

# ---------- PALETA ----------
ROSA="FFFF6B9D"; ROSA_D="FFE85A8A"; MORADO="FF8E44AD"; MORADO_L="FFA569BD"
ROSA_P="FFFADBD8"; LILA="FFF5EEF8"; CELESTE="FFAED6F1"; ROJO="FFF1948A"
BLANCO="FFFFFFFF"; TXT="FF4A2C3D"; GRIS="FFF8F9F9"
VERDE="FFD5F5E3"; VERDE_T="FF196F3D"; AMBAR="FFFDEBD0"; AMBAR_T="FFB9770E"
AMAR="FFFCF3CF"; AMAR_T="FF9A7D0A"; ROJO_T="FF7B241C"
FUENTE="Segoe UI"
AZUL_IN="FF0000FF"   # celdas de captura manual

MON='"S/ "#,##0.00'; MON0='"S/ "#,##0'; PCT='0.0%'; NUM='#,##0'; FEC='dd/mm/yyyy'

def F(sz=10,b=False,color=TXT,it=False): return Font(name=FUENTE,size=sz,bold=b,color=color,italic=it)
def P(c): return PatternFill("solid",fgColor=c)
thin=Side(style="thin",color=MORADO_L); med=Side(style="medium",color=MORADO)
BD=Border(left=thin,right=thin,top=thin,bottom=thin)
BD_CARD=Border(left=med,right=med,top=med,bottom=med)
CT=Alignment(horizontal="center",vertical="center",wrap_text=True)
LT=Alignment(horizontal="left",vertical="center")
RT=Alignment(horizontal="right",vertical="center")

wb=Workbook()

# =========================================================
# 0_Listas_y_Apoyo
# =========================================================
ws0=wb.active; ws0.title="0_Listas_y_Apoyo"
CATS=["Libros","Papelería","Arte","Kawaii/Escolar","Escritura","Cuadernos","Mochilas y Loncheras","Tecnología","Manualidades","Oficina"]
PAGOS=["Efectivo","Tarjeta","Transferencia","Yape/Plin","Crédito"]
PROVS=["Distribuidora Lima Papel SAC","Importadora Kawaii Perú EIRL","Editorial Andina SAC","Arte & Color Distribuciones","Faber Perú SA","Mayorista Mesa Redonda"]
ESTADOS=["🟢 OK","🟡 VIGILAR","🟠 REPONER","🔴 SIN STOCK"]

ws0["A1"]="PARÁMETROS Y LISTAS MAESTRAS DAMABIS"; ws0["A1"].font=F(14,True,BLANCO); ws0["A1"].fill=P(ROSA)
ws0.merge_cells("A1:F1"); ws0.row_dimensions[1].height=26; ws0["A1"].alignment=CT
for col,(t,vals) in enumerate({"Categorías":CATS,"Métodos de Pago":PAGOS,"Proveedores":PROVS,"Estados de Stock":ESTADOS}.items(),start=1):
    c=ws0.cell(row=3,column=col,value=t); c.font=F(10,True,BLANCO); c.fill=P(MORADO); c.alignment=CT; c.border=BD
    for i,v in enumerate(vals,start=4):
        cc=ws0.cell(row=i,column=col,value=v); cc.font=F(10); cc.border=BD
        cc.fill=P(LILA if i%2==0 else BLANCO)
for w,col in [(26,"A"),(18,"B"),(34,"C"),(18,"D"),(3,"E"),(14,"F")]: ws0.column_dimensions[col].width=w
ws0.sheet_properties.tabColor="A569BD"

# ---- Tabla de apoyo para Dashboard (Top productos) : filas 3..202 col H..M
ws0["H1"]="TABLA DE APOYO · RANKING DE PRODUCTOS (no editar)"; ws0["H1"].font=F(11,True,BLANCO); ws0["H1"].fill=P(MORADO); ws0.merge_cells("H1:M1"); ws0["H1"].alignment=CT
apoyo=["Código/SKU","Descripción","Unidades Mes","Ingresos Mes","Ganancia Mes","Índice Ranking"]
for j,t in enumerate(apoyo,start=8):
    c=ws0.cell(row=3,column=j,value=t); c.font=F(10,True,BLANCO); c.fill=P(ROSA_D); c.alignment=CT; c.border=BD
    ws0.column_dimensions[C(j)].width=[16,34,13,15,15,15][j-8]

# =========================================================
# Helpers de layout
# =========================================================
def banner(ws, titulo, subtitulo, ncols, logo=True):
    ws.merge_cells(start_row=1,start_column=1,end_row=4,end_column=ncols)
    c=ws.cell(row=1,column=1)
    c.value=f"   {titulo}\n   {subtitulo}"
    c.font=F(20,True,BLANCO); c.fill=P(ROSA); c.alignment=Alignment(horizontal="left",vertical="center",wrap_text=True)
    for r,h in [(1,20),(2,22),(3,22),(4,20),(5,6)]: ws.row_dimensions[r].height=h
    if logo:
        im=XLImage(LOGO); im.width=92; im.height=92
        im.anchor=f"{C(max(1,ncols-1))}1"
        ws.add_image(im)

def kpi(ws, row, col, span, titulo, formula, fmt, fill=MORADO):
    """Tarjeta KPI: título en 'row', valor en row+1..row+2"""
    ws.merge_cells(start_row=row,start_column=col,end_row=row,end_column=col+span-1)
    t=ws.cell(row=row,column=col,value=titulo); t.font=F(9,True,BLANCO); t.fill=P(fill); t.alignment=CT
    ws.merge_cells(start_row=row+1,start_column=col,end_row=row+2,end_column=col+span-1)
    v=ws.cell(row=row+1,column=col,value=formula); v.font=F(16,True,MORADO); v.fill=P(LILA); v.alignment=CT; v.number_format=fmt
    for rr in (row,row+1,row+2):
        for cc in range(col,col+span):
            ws.cell(row=rr,column=cc).border=BD_CARD
    ws.row_dimensions[row].height=18; ws.row_dimensions[row+1].height=14; ws.row_dimensions[row+2].height=16

def encabezados(ws,row,cols,widths):
    for j,(t,w) in enumerate(zip(cols,widths),start=1):
        c=ws.cell(row=row,column=j,value=t); c.font=F(10,True,BLANCO); c.fill=P(ROSA_D); c.alignment=CT; c.border=BD
        ws.column_dimensions[C(j)].width=w
    ws.row_dimensions[row].height=34

def zebra(ws,row_ini,row_fin,ncols):
    rng=f"A{row_ini}:{C(ncols)}{row_fin}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f"AND(MOD(ROW(),2)=0,$A{row_ini}<>\"\")"], fill=P(LILA), stopIfTrue=False))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f"AND(MOD(ROW(),2)=1,$A{row_ini}<>\"\")"], fill=P(ROSA_P), stopIfTrue=False))

def leyenda(ws,row,ncols,texto):
    ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=ncols)
    c=ws.cell(row=row,column=1,value=texto); c.font=F(9,False,MORADO,it=True); c.fill=P(CELESTE); c.alignment=LT
    ws.row_dimensions[row].height=16

HDR=10; INI=11
C_FIN=510      # compras: filas 11..510
S_FIN=210      # stock:   filas 11..210
V_FIN=610      # ventas:  filas 11..610

SH1="'1_Compras_e_Inventario_Inicial'"
SH2="'2_Control_de_Stock'"
SH3="'3_Ventas_Diarias'"

# =========================================================
# 1_Compras_e_Inventario_Inicial
# =========================================================
ws1=wb.create_sheet("1_Compras_e_Inventario_Inicial"); ws1.sheet_properties.tabColor="FF6B9D"
banner(ws1,"DAMABIS · Compras e Inventario Inicial","Registro de mercancía ingresada · Librería y Útiles Escolares",9)
cols1=["Fecha","ID Transacción","Proveedor","Código/SKU","Descripción del Producto","Categoría","Cantidad Comprada","Costo Unitario","Costo Total"]
w1=[12,16,30,14,38,20,13,14,15]
kpi(ws1,6,1,2,"💰 GASTO TOTAL EN COMPRAS",f"=SUM(I{INI}:I{C_FIN})",MON)
kpi(ws1,6,3,2,"📦 UNIDADES INGRESADAS",f"=SUM(G{INI}:G{C_FIN})",NUM)
kpi(ws1,6,5,2,"🧾 N° DE TRANSACCIONES",f'=SUMPRODUCT((B{INI}:B{C_FIN}<>"")/COUNTIF(B{INI}:B{C_FIN},B{INI}:B{C_FIN}&""))',NUM)
kpi(ws1,6,7,1,"🏭 PROVEEDORES",f'=SUMPRODUCT((C{INI}:C{C_FIN}<>"")/COUNTIF(C{INI}:C{C_FIN},C{INI}:C{C_FIN}&""))',NUM)
kpi(ws1,6,8,2,"📊 COSTO UNIT. PROMEDIO",f"=IFERROR(SUM(I{INI}:I{C_FIN})/SUM(G{INI}:G{C_FIN}),0)",MON)
leyenda(ws1,9,9,"   ✏️ Celdas en AZUL = captura manual.  Celdas en NEGRO = fórmula automática (no escribir).  Usa las listas desplegables en Proveedor, SKU y Categoría.")
encabezados(ws1,HDR,cols1,w1)

compras=[
 ("2026-08-02","C-2026-001","Distribuidora Lima Papel SAC","PAP-001","Papel Bond A4 75g (millar)","Papelería",30,13.50),
 ("2026-08-02","C-2026-001","Distribuidora Lima Papel SAC","CUA-010","Cuaderno A4 cuadriculado 100h","Cuadernos",120,4.20),
 ("2026-08-03","C-2026-002","Importadora Kawaii Perú EIRL","KAW-101","Set stickers ajolote kawaii","Kawaii/Escolar",200,2.10),
 ("2026-08-03","C-2026-002","Importadora Kawaii Perú EIRL","KAW-102","Lapicero gel pastel x6","Kawaii/Escolar",150,5.40),
 ("2026-08-05","C-2026-003","Editorial Andina SAC","LIB-201","Libro infantil ilustrado","Libros",60,18.00),
 ("2026-08-05","C-2026-003","Editorial Andina SAC","LIB-202","Diccionario escolar","Libros",25,22.00),
 ("2026-08-08","C-2026-004","Arte & Color Distribuciones","ART-301","Témpera 12 colores","Arte",80,7.80),
 ("2026-08-08","C-2026-004","Arte & Color Distribuciones","ART-302","Block de dibujo A3","Arte",90,6.30),
 ("2026-08-10","C-2026-005","Faber Perú SA","ESC-401","Colores x24 largos","Escritura",100,11.50),
 ("2026-08-10","C-2026-005","Faber Perú SA","ESC-402","Plumón resaltador pastel x4","Escritura",140,4.90),
 ("2026-08-12","C-2026-006","Mayorista Mesa Redonda","MOC-501","Mochila escolar estampada","Mochilas y Loncheras",8,32.00),
 ("2026-08-12","C-2026-006","Mayorista Mesa Redonda","MOC-502","Lonchera térmica kawaii","Mochilas y Loncheras",10,21.00),
 ("2026-08-15","C-2026-007","Distribuidora Lima Papel SAC","PAP-002","Cartulina de colores (pliego)","Papelería",300,0.80),
 ("2026-08-15","C-2026-007","Distribuidora Lima Papel SAC","OFI-601","Archivador palanca A4","Oficina",60,9.20),
 ("2026-08-18","C-2026-008","Importadora Kawaii Perú EIRL","KAW-103","Peluche llavero ajolote","Kawaii/Escolar",90,8.60),
 ("2026-08-20","C-2026-009","Arte & Color Distribuciones","MAN-701","Set de foamy glitter x10","Manualidades",70,5.70),
 ("2026-08-22","C-2026-010","Mayorista Mesa Redonda","TEC-801","Calculadora científica","Tecnología",4,27.00),
 ("2026-08-25","C-2026-011","Faber Perú SA","ESC-403","Lápiz 2B caja x12","Escritura",110,6.40),
 ("2026-08-28","C-2026-012","Editorial Andina SAC","LIB-203","Plan lector secundaria","Libros",45,16.50),
 ("2026-09-01","C-2026-013","Distribuidora Lima Papel SAC","CUA-011","Cuaderno anillado A5 80h","Cuadernos",130,3.90),
]
for i,r in enumerate(compras):
    row=INI+i
    ws1.cell(row=row,column=1,value=dt.date.fromisoformat(r[0]))
    for j,v in enumerate(r[1:],start=2): ws1.cell(row=row,column=j,value=v)

for row in range(INI,C_FIN+1):
    ws1.cell(row=row,column=9,value=f'=IF(OR($G{row}="",$H{row}=""),"",$G{row}*$H{row})')
    for j in range(1,10):
        c=ws1.cell(row=row,column=j); c.border=BD
        c.font=F(10,color=(AZUL_IN if j<=8 else TXT))
    ws1.cell(row=row,column=1).number_format=FEC; ws1.cell(row=row,column=1).alignment=CT
    ws1.cell(row=row,column=2).alignment=CT; ws1.cell(row=row,column=4).alignment=CT
    ws1.cell(row=row,column=6).alignment=CT
    ws1.cell(row=row,column=7).number_format=NUM; ws1.cell(row=row,column=7).alignment=CT
    ws1.cell(row=row,column=8).number_format=MON
    ws1.cell(row=row,column=9).number_format=MON; ws1.cell(row=row,column=9).font=F(10,True,TXT)
zebra(ws1,INI,C_FIN,9)
ws1.auto_filter.ref=f"A{HDR}:I{C_FIN}"; ws1.freeze_panes=f"A{INI}"
ws1["A10"].comment=Comment("Supuesto: fechas de ejemplo ago-sep 2026. Reemplaza las 20 filas de muestra por tus compras reales.","DAMABIS")

# =========================================================
# 2_Control_de_Stock
# =========================================================
ws2=wb.create_sheet("2_Control_de_Stock"); ws2.sheet_properties.tabColor="8E44AD"
banner(ws2,"DAMABIS · Control de Stock","Maestro de productos · Cálculo automático de existencias y márgenes",14)
cols2=["Código/SKU","Descripción","Categoría","Stock Inicial","Entradas Totales","Salidas Totales","Stock Actual","Stock Mínimo","Alerta de Estado","Costo Promedio Unit.","Precio de Venta","Valor Total a Costo","Valor Total a Venta","% Margen Bruto"]
w2=[14,36,20,11,11,11,11,11,18,14,13,15,15,12]
kpi(ws2,6,1,2,"🧮 SKUs ACTIVOS",f'=COUNTIF(A{INI}:A{S_FIN},"?*")',NUM)
kpi(ws2,6,3,2,"🏦 VALOR INVENTARIO (COSTO)",f"=SUM(L{INI}:L{S_FIN})",MON)
kpi(ws2,6,5,2,"💵 VALOR INVENTARIO (VENTA)",f"=SUM(M{INI}:M{S_FIN})",MON)
kpi(ws2,6,7,2,"📈 MARGEN PROMEDIO",f"=IFERROR((SUM(M{INI}:M{S_FIN})-SUM(L{INI}:L{S_FIN}))/SUM(M{INI}:M{S_FIN}),0)",PCT)
kpi(ws2,6,9,2,"🔴 SKUs SIN STOCK",f'=COUNTIF(I{INI}:I{S_FIN},"*SIN STOCK*")',NUM,ROJO)
kpi(ws2,6,11,2,"🟠 SKUs POR REPONER",f'=COUNTIF(I{INI}:I{S_FIN},"*REPONER*")',NUM,ROSA_D)
kpi(ws2,6,13,2,"📦 UNIDADES EN TIENDA",f"=SUM(G{INI}:G{S_FIN})",NUM)
leyenda(ws2,9,14,"   ✏️ Captura manual (azul): SKU · Descripción · Categoría · Stock Inicial · Stock Mínimo · Precio de Venta.   ⚙️ El resto se calcula solo desde las hojas 1 y 3.")
encabezados(ws2,HDR,cols2,w2)

maestro=[
 ("PAP-001","Papel Bond A4 75g (millar)","Papelería",5,4,19.90),
 ("PAP-002","Cartulina de colores (pliego)","Papelería",50,60,1.50),
 ("CUA-010","Cuaderno A4 cuadriculado 100h","Cuadernos",20,30,6.50),
 ("CUA-011","Cuaderno anillado A5 80h","Cuadernos",15,25,5.90),
 ("KAW-101","Set stickers ajolote kawaii","Kawaii/Escolar",30,40,4.50),
 ("KAW-102","Lapicero gel pastel x6","Kawaii/Escolar",25,30,9.90),
 ("KAW-103","Peluche llavero ajolote","Kawaii/Escolar",10,20,16.90),
 ("LIB-201","Libro infantil ilustrado","Libros",8,12,29.90),
 ("LIB-202","Diccionario escolar","Libros",5,8,34.90),
 ("LIB-203","Plan lector secundaria","Libros",6,10,26.90),
 ("ART-301","Témpera 12 colores","Arte",12,20,13.90),
 ("ART-302","Block de dibujo A3","Arte",15,20,10.90),
 ("ESC-401","Colores x24 largos","Escritura",18,25,19.90),
 ("ESC-402","Plumón resaltador pastel x4","Escritura",22,30,8.90),
 ("ESC-403","Lápiz 2B caja x12","Escritura",20,25,10.90),
 ("MOC-501","Mochila escolar estampada","Mochilas y Loncheras",4,8,59.90),
 ("MOC-502","Lonchera térmica kawaii","Mochilas y Loncheras",6,10,38.90),
 ("OFI-601","Archivador palanca A4","Oficina",10,15,15.90),
 ("MAN-701","Set de foamy glitter x10","Manualidades",8,15,9.90),
 ("TEC-801","Calculadora científica","Tecnología",3,6,45.00),
 ("KAW-104","Cuaderno ajolote tapa dura A5","Kawaii/Escolar",0,12,22.90),
]
for i,r in enumerate(maestro):
    row=INI+i
    ws2.cell(row=row,column=1,value=r[0]); ws2.cell(row=row,column=2,value=r[1]); ws2.cell(row=row,column=3,value=r[2])
    ws2.cell(row=row,column=4,value=r[3]); ws2.cell(row=row,column=8,value=r[4]); ws2.cell(row=row,column=11,value=r[5])

for row in range(INI,S_FIN+1):
    a=f"$A{row}"
    ws2.cell(row=row,column=5,value=f'=IF({a}="","",SUMIFS({SH1}!$G${INI}:$G${C_FIN},{SH1}!$D${INI}:$D${C_FIN},{a}))')
    ws2.cell(row=row,column=6,value=f'=IF({a}="","",SUMIFS({SH3}!$E${INI}:$E${V_FIN},{SH3}!$C${INI}:$C${V_FIN},{a}))')
    ws2.cell(row=row,column=7,value=f'=IF({a}="","",N($D{row})+N($E{row})-N($F{row}))')
    ws2.cell(row=row,column=9,value=(f'=IF({a}="","",IF($G{row}<=0,"🔴 SIN STOCK",'
                                     f'IF($G{row}<=$H{row},"🟠 REPONER",'
                                     f'IF($G{row}<=$H{row}*1.5,"🟡 VIGILAR","🟢 OK"))))'))
    ws2.cell(row=row,column=10,value=(f'=IF({a}="","",IFERROR(SUMIFS({SH1}!$I${INI}:$I${C_FIN},{SH1}!$D${INI}:$D${C_FIN},{a})'
                                      f'/SUMIFS({SH1}!$G${INI}:$G${C_FIN},{SH1}!$D${INI}:$D${C_FIN},{a}),0))'))
    ws2.cell(row=row,column=12,value=f'=IF({a}="","",$G{row}*$J{row})')
    ws2.cell(row=row,column=13,value=f'=IF({a}="","",$G{row}*$K{row})')
    ws2.cell(row=row,column=14,value=f'=IF({a}="","",IFERROR(($K{row}-$J{row})/$K{row},0))')
    for j in range(1,15):
        c=ws2.cell(row=row,column=j); c.border=BD
        c.font=F(10,color=(AZUL_IN if j in (1,2,3,4,8,11) else TXT))
    for j in (4,5,6,8): ws2.cell(row=row,column=j).number_format=NUM; ws2.cell(row=row,column=j).alignment=CT
    ws2.cell(row=row,column=1).alignment=CT; ws2.cell(row=row,column=3).alignment=CT
    ws2.cell(row=row,column=7).number_format=NUM; ws2.cell(row=row,column=7).alignment=CT; ws2.cell(row=row,column=7).font=F(11,True,MORADO)
    ws2.cell(row=row,column=9).alignment=CT; ws2.cell(row=row,column=9).font=F(10,True)
    for j in (10,11,12,13): ws2.cell(row=row,column=j).number_format=MON
    ws2.cell(row=row,column=14).number_format=PCT; ws2.cell(row=row,column=14).alignment=CT
zebra(ws2,INI,S_FIN,14)
rngI=f"I{INI}:I{S_FIN}"
ws2.conditional_formatting.add(rngI,FormulaRule(formula=[f'ISNUMBER(SEARCH("SIN STOCK",$I{INI}))'],fill=P(ROJO),font=Font(name=FUENTE,size=10,bold=True,color=ROJO_T),stopIfTrue=True))
ws2.conditional_formatting.add(rngI,FormulaRule(formula=[f'ISNUMBER(SEARCH("REPONER",$I{INI}))'],fill=P(AMBAR),font=Font(name=FUENTE,size=10,bold=True,color=AMBAR_T),stopIfTrue=True))
ws2.conditional_formatting.add(rngI,FormulaRule(formula=[f'ISNUMBER(SEARCH("VIGILAR",$I{INI}))'],fill=P(AMAR),font=Font(name=FUENTE,size=10,bold=True,color=AMAR_T),stopIfTrue=True))
ws2.conditional_formatting.add(rngI,FormulaRule(formula=[f'ISNUMBER(SEARCH("OK",$I{INI}))'],fill=P(VERDE),font=Font(name=FUENTE,size=10,bold=True,color=VERDE_T),stopIfTrue=True))
# fila completa resaltada si está bajo mínimo
ws2.conditional_formatting.add(f"A{INI}:N{S_FIN}",FormulaRule(formula=[f'AND($A{INI}<>"",$G{INI}<=$H{INI})'],fill=P("FFFDECEA"),stopIfTrue=False))
ws2.conditional_formatting.add(f"G{INI}:G{S_FIN}",DataBarRule(start_type="num",start_value=0,end_type="max",color="8E44AD",showValue=True))
ws2.conditional_formatting.add(f"N{INI}:N{S_FIN}",ColorScaleRule(start_type="num",start_value=0,start_color="FFF1948A",mid_type="num",mid_value=0.35,mid_color="FFFCF3CF",end_type="num",end_value=0.7,end_color="FFD5F5E3"))
ws2.auto_filter.ref=f"A{HDR}:N{S_FIN}"; ws2.freeze_panes=f"D{INI}"

# =========================================================
# 3_Ventas_Diarias
# =========================================================
ws3=wb.create_sheet("3_Ventas_Diarias"); ws3.sheet_properties.tabColor="E85A8A"
banner(ws3,"DAMABIS · Ventas Diarias","Registro de ticket por ticket · Costo y ganancia automáticos",11)
cols3=["Fecha","N° Ticket/Comprobante","Código/SKU","Descripción","Cantidad Vendida","Precio Venta Unitario","Total Cobrado","Método de Pago","Costo de lo Vendido","Ganancia Bruta","Categoría (aux.)"]
w3=[12,18,14,36,11,14,14,16,15,14,18]
kpi(ws3,6,1,2,"🛍️ VENTAS ACUMULADAS",f"=SUM(G{INI}:G{V_FIN})",MON)
kpi(ws3,6,3,2,"💚 GANANCIA BRUTA",f"=SUM(J{INI}:J{V_FIN})",MON)
kpi(ws3,6,5,2,"📦 UNIDADES VENDIDAS",f"=SUM(E{INI}:E{V_FIN})",NUM)
kpi(ws3,6,7,2,"🎟️ TICKET PROMEDIO",f'=IFERROR(SUM(G{INI}:G{V_FIN})/SUMPRODUCT((B{INI}:B{V_FIN}<>"")/COUNTIF(B{INI}:B{V_FIN},B{INI}:B{V_FIN}&"")),0)',MON)
kpi(ws3,6,9,3,"📊 MARGEN SOBRE VENTAS",f"=IFERROR(SUM(J{INI}:J{V_FIN})/SUM(G{INI}:G{V_FIN}),0)",PCT)
leyenda(ws3,9,11,"   ✏️ Captura manual (azul): Fecha · N° Ticket · SKU · Cantidad · Método de Pago.   ⚙️ Descripción, Precio, Total, COGS, Ganancia y Categoría se llenan solos con BUSCARV.")
encabezados(ws3,HDR,cols3,w3)

ventas=[
 ("2026-08-05","T-0001","CUA-010",12,"Efectivo"),("2026-08-05","T-0002","KAW-101",8,"Yape/Plin"),
 ("2026-08-06","T-0003","ESC-401",5,"Tarjeta"),("2026-08-06","T-0004","PAP-001",3,"Efectivo"),
 ("2026-08-07","T-0005","LIB-201",4,"Transferencia"),("2026-08-08","T-0006","KAW-102",10,"Efectivo"),
 ("2026-08-09","T-0007","ART-301",6,"Yape/Plin"),("2026-08-10","T-0008","MOC-501",2,"Tarjeta"),
 ("2026-08-11","T-0009","CUA-011",14,"Efectivo"),("2026-08-12","T-0010","ESC-402",18,"Efectivo"),
 ("2026-08-13","T-0011","KAW-103",7,"Yape/Plin"),("2026-08-14","T-0012","LIB-202",3,"Tarjeta"),
 ("2026-08-15","T-0013","PAP-002",40,"Efectivo"),("2026-08-16","T-0014","ART-302",9,"Transferencia"),
 ("2026-08-17","T-0015","OFI-601",5,"Tarjeta"),("2026-08-18","T-0016","KAW-101",15,"Efectivo"),
 ("2026-08-19","T-0017","ESC-403",11,"Yape/Plin"),("2026-08-20","T-0018","MAN-701",8,"Efectivo"),
 ("2026-08-21","T-0019","TEC-801",2,"Tarjeta"),("2026-08-22","T-0020","MOC-502",4,"Transferencia"),
 ("2026-08-23","T-0021","CUA-010",20,"Efectivo"),("2026-08-24","T-0022","KAW-102",12,"Yape/Plin"),
 ("2026-08-25","T-0023","LIB-203",6,"Tarjeta"),("2026-08-26","T-0024","ESC-401",9,"Efectivo"),
 ("2026-08-27","T-0025","ART-301",7,"Efectivo"),("2026-08-28","T-0026","PAP-001",4,"Transferencia"),
 ("2026-08-29","T-0027","KAW-103",6,"Yape/Plin"),("2026-08-30","T-0028","CUA-011",16,"Efectivo"),
 ("2026-08-31","T-0029","MOC-501",3,"Tarjeta"),("2026-08-31","T-0030","LIB-201",5,"Efectivo"),
]
for i,r in enumerate(ventas):
    row=INI+i
    ws3.cell(row=row,column=1,value=dt.date.fromisoformat(r[0]))
    ws3.cell(row=row,column=2,value=r[1]); ws3.cell(row=row,column=3,value=r[2])
    ws3.cell(row=row,column=5,value=r[3]); ws3.cell(row=row,column=8,value=r[4])

MAESTRO=f"{SH2}!$A${INI}:$N${S_FIN}"
for row in range(INI,V_FIN+1):
    c3=f"$C{row}"
    ws3.cell(row=row,column=4,value=f'=IF({c3}="","",IFERROR(VLOOKUP({c3},{MAESTRO},2,0),"⚠️ SKU no existe"))')
    ws3.cell(row=row,column=6,value=f'=IF({c3}="","",IFERROR(VLOOKUP({c3},{MAESTRO},11,0),0))')
    ws3.cell(row=row,column=7,value=f'=IF(OR({c3}="",$E{row}=""),"",$E{row}*$F{row})')
    ws3.cell(row=row,column=9,value=f'=IF({c3}="","",IFERROR($E{row}*VLOOKUP({c3},{MAESTRO},10,0),0))')
    ws3.cell(row=row,column=10,value=f'=IF($G{row}="","",$G{row}-$I{row})')
    ws3.cell(row=row,column=11,value=f'=IF({c3}="","",IFERROR(VLOOKUP({c3},{MAESTRO},3,0),""))')
    for j in range(1,12):
        c=ws3.cell(row=row,column=j); c.border=BD
        c.font=F(10,color=(AZUL_IN if j in (1,2,3,5,8) else TXT))
    ws3.cell(row=row,column=1).number_format=FEC; ws3.cell(row=row,column=1).alignment=CT
    ws3.cell(row=row,column=2).alignment=CT; ws3.cell(row=row,column=3).alignment=CT
    ws3.cell(row=row,column=5).number_format=NUM; ws3.cell(row=row,column=5).alignment=CT
    for j in (6,7,9,10): ws3.cell(row=row,column=j).number_format=MON
    ws3.cell(row=row,column=7).font=F(10,True,TXT)
    ws3.cell(row=row,column=8).alignment=CT
    ws3.cell(row=row,column=10).font=F(10,True,VERDE_T)
    ws3.cell(row=row,column=11).alignment=CT; ws3.cell(row=row,column=11).font=F(9,color="FF7D6B76",it=True)
zebra(ws3,INI,V_FIN,11)
ws3.conditional_formatting.add(f"J{INI}:J{V_FIN}",FormulaRule(formula=[f'AND($G{INI}<>"",$J{INI}<0)'],fill=P(ROJO),font=Font(name=FUENTE,size=10,bold=True,color=ROJO_T)))
ws3.conditional_formatting.add(f"D{INI}:D{V_FIN}",FormulaRule(formula=[f'ISNUMBER(SEARCH("no existe",$D{INI}))'],fill=P(ROJO),font=Font(name=FUENTE,size=10,bold=True,color=ROJO_T)))
ws3.auto_filter.ref=f"A{HDR}:K{V_FIN}"; ws3.freeze_panes=f"D{INI}"

# =========================================================
# Tabla de apoyo (0_Listas) alimentada por mes del Dashboard
# =========================================================
SH4="'4_Dashboard_DAMABIS'"
for i in range(200):
    row=4+i; sku_row=INI+i
    ws0.cell(row=row,column=8,value=f'=IF({SH2}!$A${sku_row}="","",{SH2}!$A${sku_row})')
    ws0.cell(row=row,column=9,value=f'=IF({SH2}!$A${sku_row}="","",{SH2}!$B${sku_row})')
    ws0.cell(row=row,column=10,value=f'=IF($H{row}="","",SUMIFS({SH3}!$E${INI}:$E${V_FIN},{SH3}!$C${INI}:$C${V_FIN},$H{row},{SH3}!$A${INI}:$A${V_FIN},">="&{SH4}!$P$5,{SH3}!$A${INI}:$A${V_FIN},"<="&{SH4}!$P$6))')
    ws0.cell(row=row,column=11,value=f'=IF($H{row}="","",SUMIFS({SH3}!$G${INI}:$G${V_FIN},{SH3}!$C${INI}:$C${V_FIN},$H{row},{SH3}!$A${INI}:$A${V_FIN},">="&{SH4}!$P$5,{SH3}!$A${INI}:$A${V_FIN},"<="&{SH4}!$P$6))')
    ws0.cell(row=row,column=12,value=f'=IF($H{row}="","",SUMIFS({SH3}!$J${INI}:$J${V_FIN},{SH3}!$C${INI}:$C${V_FIN},$H{row},{SH3}!$A${INI}:$A${V_FIN},">="&{SH4}!$P$5,{SH3}!$A${INI}:$A${V_FIN},"<="&{SH4}!$P$6))')
    ws0.cell(row=row,column=13,value=f'=IF($H{row}="","",N($K{row})+ROW()/1000000)')
    for j in range(8,14):
        c=ws0.cell(row=row,column=j); c.border=BD; c.font=F(9)
        c.fill=P(LILA if row%2==0 else BLANCO)
    ws0.cell(row=row,column=10).number_format=NUM
    for j in (11,12): ws0.cell(row=row,column=j).number_format=MON
    ws0.cell(row=row,column=13).number_format="0.000000"

# =========================================================
# 4_Dashboard_DAMABIS
# =========================================================
ws4=wb.create_sheet("4_Dashboard_DAMABIS"); ws4.sheet_properties.tabColor="FF6B9D"
ws4.sheet_view.showGridLines=False
banner(ws4,"DAMABIS ✨ Panel de Control Gerencial","Librería y Útiles Escolares · Indicadores del mes seleccionado",14)
for j,w in enumerate([15,13,15,13,15,13,15,13,15,13,15,13,15,13,3,14,14],start=1):
    ws4.column_dimensions[C(j)].width=w

# Selector de periodo
ws4.merge_cells("A6:B7"); s=ws4["A6"]; s.value="📅 PERIODO"; s.font=F(11,True,BLANCO); s.fill=P(MORADO); s.alignment=CT; 
for r in (6,7):
    for c in (1,2): ws4.cell(row=r,column=c).border=BD_CARD
ws4["C6"]="Mes"; ws4["C7"]="Año"
for a in ("C6","C7"): ws4[a].font=F(10,True,MORADO); ws4[a].alignment=RT; ws4[a].fill=P(CELESTE); ws4[a].border=BD
ws4["D6"]=8; ws4["D7"]=2026
for a in ("D6","D7"):
    ws4[a].font=F(12,True,AZUL_IN); ws4[a].fill=P("FFFFFF00"); ws4[a].alignment=CT; ws4[a].border=BD_CARD
ws4["E6"]='=CHOOSE($D$6,"ENERO","FEBRERO","MARZO","ABRIL","MAYO","JUNIO","JULIO","AGOSTO","SEPTIEMBRE","OCTUBRE","NOVIEMBRE","DICIEMBRE")&" "&$D$7'
ws4["E6"].font=F(12,True,ROSA_D); ws4["E6"].alignment=LT
ws4.merge_cells("E6:H7")
ws4["P4"]="Rango calculado (auxiliar)"; ws4["P4"].font=F(9,True,MORADO)
ws4["P5"]="=DATE($D$7,$D$6,1)"; ws4["P6"]="=EOMONTH($P$5,0)"
for a in ("P5","P6"): ws4[a].number_format=FEC; ws4[a].font=F(9)
INIm="$P$5"; FINm="$P$6"

# KPI cards principales
kpi(ws4,9,1,3,"🛍️ VENTAS TOTALES DEL MES",
    f'=SUMIFS({SH3}!$G${INI}:$G${V_FIN},{SH3}!$A${INI}:$A${V_FIN},">="&{INIm},{SH3}!$A${INI}:$A${V_FIN},"<="&{FINm})',MON,ROSA)
kpi(ws4,9,4,3,"💚 GANANCIA BRUTA DEL MES",
    f'=SUMIFS({SH3}!$J${INI}:$J${V_FIN},{SH3}!$A${INI}:$A${V_FIN},">="&{INIm},{SH3}!$A${INI}:$A${V_FIN},"<="&{FINm})',MON,MORADO)
kpi(ws4,9,7,3,"💰 INVERSIÓN EN COMPRAS DEL MES",
    f'=SUMIFS({SH1}!$I${INI}:$I${C_FIN},{SH1}!$A${INI}:$A${C_FIN},">="&{INIm},{SH1}!$A${INI}:$A${C_FIN},"<="&{FINm})',MON,ROSA_D)
kpi(ws4,9,10,3,"🏦 VALOR TOTAL DEL INVENTARIO",f"=SUM({SH2}!$L${INI}:$L${S_FIN})",MON,MORADO_L)
kpi(ws4,9,13,2,"📈 MARGEN DEL MES",
    f'=IFERROR(SUMIFS({SH3}!$J${INI}:$J${V_FIN},{SH3}!$A${INI}:$A${V_FIN},">="&{INIm},{SH3}!$A${INI}:$A${V_FIN},"<="&{FINm})'
    f'/SUMIFS({SH3}!$G${INI}:$G${V_FIN},{SH3}!$A${INI}:$A${V_FIN},">="&{INIm},{SH3}!$A${INI}:$A${V_FIN},"<="&{FINm}),0)',PCT,ROSA)

# Fila 2 de KPIs
kpi(ws4,13,1,3,"⭐ PRODUCTO ESTRELLA DEL MES",
    f'=IFERROR(INDEX(\'0_Listas_y_Apoyo\'!$I$4:$I$203,MATCH(MAX(\'0_Listas_y_Apoyo\'!$M$4:$M$203),\'0_Listas_y_Apoyo\'!$M$4:$M$203,0)),"Sin ventas")',"General",MORADO)
ws4["A14"].font=F(11,True,MORADO)
kpi(ws4,13,4,3,"🎟️ TICKETS DEL MES",
    f'=SUMPRODUCT(({SH3}!$B${INI}:$B${V_FIN}<>"")*({SH3}!$A${INI}:$A${V_FIN}>={INIm})*({SH3}!$A${INI}:$A${V_FIN}<={FINm})'
    f'/COUNTIF({SH3}!$B${INI}:$B${V_FIN},{SH3}!$B${INI}:$B${V_FIN}&""))',NUM,ROSA_D)
kpi(ws4,13,7,3,"📦 UNIDADES VENDIDAS EN EL MES",
    f'=SUMIFS({SH3}!$E${INI}:$E${V_FIN},{SH3}!$A${INI}:$A${V_FIN},">="&{INIm},{SH3}!$A${INI}:$A${V_FIN},"<="&{FINm})',NUM,MORADO_L)
kpi(ws4,13,10,3,"🚨 SKUs EN ALERTA",
    f'=COUNTIF({SH2}!$I${INI}:$I${S_FIN},"*REPONER*")+COUNTIF({SH2}!$I${INI}:$I${S_FIN},"*SIN STOCK*")',NUM,ROJO)
kpi(ws4,13,13,2,"💳 % VENTAS EN EFECTIVO",
    f'=IFERROR(SUMIFS({SH3}!$G${INI}:$G${V_FIN},{SH3}!$H${INI}:$H${V_FIN},"Efectivo",{SH3}!$A${INI}:$A${V_FIN},">="&{INIm},{SH3}!$A${INI}:$A${V_FIN},"<="&{FINm})'
    f'/SUMIFS({SH3}!$G${INI}:$G${V_FIN},{SH3}!$A${INI}:$A${V_FIN},">="&{INIm},{SH3}!$A${INI}:$A${V_FIN},"<="&{FINm}),0)',PCT,ROSA)

# --- Tabla: Ventas por Categoría (A18)
r0=18
ws4.merge_cells(f"A{r0}:D{r0}"); t=ws4[f"A{r0}"]; t.value="📚 VENTAS POR CATEGORÍA"; t.font=F(12,True,BLANCO); t.fill=P(MORADO); t.alignment=CT
hdrs=["Categoría","Unidades","Ingresos","% Part."]
for j,h in enumerate(hdrs,start=1):
    c=ws4.cell(row=r0+1,column=j,value=h); c.font=F(10,True,BLANCO); c.fill=P(ROSA_D); c.alignment=CT; c.border=BD
for i,cat in enumerate(CATS):
    row=r0+2+i
    ws4.cell(row=row,column=1,value=f"='0_Listas_y_Apoyo'!$A${4+i}")
    ws4.cell(row=row,column=2,value=f'=SUMIFS({SH3}!$E${INI}:$E${V_FIN},{SH3}!$K${INI}:$K${V_FIN},$A{row},{SH3}!$A${INI}:$A${V_FIN},">="&{INIm},{SH3}!$A${INI}:$A${V_FIN},"<="&{FINm})')
    ws4.cell(row=row,column=3,value=f'=SUMIFS({SH3}!$G${INI}:$G${V_FIN},{SH3}!$K${INI}:$K${V_FIN},$A{row},{SH3}!$A${INI}:$A${V_FIN},">="&{INIm},{SH3}!$A${INI}:$A${V_FIN},"<="&{FINm})')
    ws4.cell(row=row,column=4,value=f"=IFERROR($C{row}/SUM($C${r0+2}:$C${r0+1+len(CATS)}),0)")
    for j in range(1,5):
        c=ws4.cell(row=row,column=j); c.border=BD; c.font=F(10); c.fill=P(LILA if i%2==0 else BLANCO)
    ws4.cell(row=row,column=2).number_format=NUM; ws4.cell(row=row,column=2).alignment=CT
    ws4.cell(row=row,column=3).number_format=MON
    ws4.cell(row=row,column=4).number_format=PCT; ws4.cell(row=row,column=4).alignment=CT
rt=r0+2+len(CATS)
ws4.cell(row=rt,column=1,value="TOTAL").font=F(10,True,BLANCO)
ws4.cell(row=rt,column=2,value=f"=SUM($B${r0+2}:$B${rt-1})")
ws4.cell(row=rt,column=3,value=f"=SUM($C${r0+2}:$C${rt-1})")
ws4.cell(row=rt,column=4,value=f"=IFERROR($C{rt}/$C{rt},0)")
for j in range(1,5):
    c=ws4.cell(row=rt,column=j); c.fill=P(MORADO); c.font=F(10,True,BLANCO); c.border=BD; c.alignment=CT if j!=1 else LT
ws4.cell(row=rt,column=3).number_format=MON; ws4.cell(row=rt,column=4).number_format=PCT
ws4.conditional_formatting.add(f"C{r0+2}:C{rt-1}",DataBarRule(start_type="num",start_value=0,end_type="max",color="FF6B9D",showValue=True))

# --- Tabla: Top 5 Productos (F18)
ws4.merge_cells(f"F{r0}:J{r0}"); t=ws4[f"F{r0}"]; t.value="🏆 TOP 5 PRODUCTOS MÁS VENDIDOS DEL MES (por ingresos S/)"; t.font=F(12,True,BLANCO); t.fill=P(MORADO); t.alignment=CT
hdrs2=["#","Descripción","SKU","Unidades","Ingresos"]
for j,h in enumerate(hdrs2,start=6):
    c=ws4.cell(row=r0+1,column=j,value=h); c.font=F(10,True,BLANCO); c.fill=P(ROSA_D); c.alignment=CT; c.border=BD
AP="'0_Listas_y_Apoyo'"
for k in range(5):
    row=r0+2+k
    big=f"LARGE({AP}!$M$4:$M$203,{k+1})"
    ws4.cell(row=row,column=6,value=k+1)
    ws4.cell(row=row,column=7,value=f'=IFERROR(IF({big}<=0,"—",INDEX({AP}!$I$4:$I$203,MATCH({big},{AP}!$M$4:$M$203,0))),"—")')
    ws4.cell(row=row,column=8,value=f'=IFERROR(IF({big}<=0,"—",INDEX({AP}!$H$4:$H$203,MATCH({big},{AP}!$M$4:$M$203,0))),"—")')
    ws4.cell(row=row,column=9,value=f'=IFERROR(IF({big}<=0,0,INDEX({AP}!$J$4:$J$203,MATCH({big},{AP}!$M$4:$M$203,0))),0)')
    ws4.cell(row=row,column=10,value=f'=IFERROR(IF({big}<=0,0,INDEX({AP}!$K$4:$K$203,MATCH({big},{AP}!$M$4:$M$203,0))),0)')
    for j in range(6,11):
        c=ws4.cell(row=row,column=j); c.border=BD; c.font=F(10); c.fill=P(LILA if k%2==0 else BLANCO)
    ws4.cell(row=row,column=6).alignment=CT; ws4.cell(row=row,column=6).font=F(11,True,MORADO)
    ws4.cell(row=row,column=8).alignment=CT
    ws4.cell(row=row,column=9).number_format=NUM; ws4.cell(row=row,column=9).alignment=CT
    ws4.cell(row=row,column=10).number_format=MON
ws4.conditional_formatting.add(f"J{r0+2}:J{r0+6}",DataBarRule(start_type="num",start_value=0,end_type="max",color="8E44AD",showValue=True))

# --- Tabla: Métodos de pago (F26)
r1=r0+8
ws4.merge_cells(f"F{r1}:J{r1}"); t=ws4[f"F{r1}"]; t.value="💳 VENTAS POR MÉTODO DE PAGO"; t.font=F(12,True,BLANCO); t.fill=P(MORADO); t.alignment=CT
for j,h in enumerate(["Método","Importe","% Part.","",""],start=6):
    c=ws4.cell(row=r1+1,column=j,value=h); c.font=F(10,True,BLANCO); c.fill=P(ROSA_D); c.alignment=CT; c.border=BD
for i,mp in enumerate(PAGOS):
    row=r1+2+i
    ws4.cell(row=row,column=6,value=f"='0_Listas_y_Apoyo'!$B${4+i}")
    ws4.cell(row=row,column=7,value=f'=SUMIFS({SH3}!$G${INI}:$G${V_FIN},{SH3}!$H${INI}:$H${V_FIN},$F{row},{SH3}!$A${INI}:$A${V_FIN},">="&{INIm},{SH3}!$A${INI}:$A${V_FIN},"<="&{FINm})')
    ws4.cell(row=row,column=8,value=f"=IFERROR($G{row}/SUM($G${r1+2}:$G${r1+1+len(PAGOS)}),0)")
    for j in range(6,9):
        c=ws4.cell(row=row,column=j); c.border=BD; c.font=F(10); c.fill=P(LILA if i%2==0 else BLANCO)
    ws4.cell(row=row,column=7).number_format=MON; ws4.cell(row=row,column=8).number_format=PCT; ws4.cell(row=row,column=8).alignment=CT

# --- Gráficos
ch1=BarChart(); ch1.type="bar"; ch1.style=10; ch1.title="Ingresos por Categoría (S/)"
d=Reference(ws4,min_col=3,min_row=r0+1,max_row=rt-1); cat=Reference(ws4,min_col=1,min_row=r0+2,max_row=rt-1)
ch1.add_data(d,titles_from_data=True); ch1.set_categories(cat); ch1.height=8.5; ch1.width=14; ch1.legend=None
try: ch1.series[0].graphicalProperties.solidFill="FF6B9D"
except Exception: pass
ws4.add_chart(ch1,f"A{rt+3}")

ch2=BarChart(); ch2.type="col"; ch2.style=10; ch2.title="Top 5 Productos · Ingresos (S/)"
d2=Reference(ws4,min_col=10,min_row=r0+1,max_row=r0+6); cat2=Reference(ws4,min_col=7,min_row=r0+2,max_row=r0+6)
ch2.add_data(d2,titles_from_data=True); ch2.set_categories(cat2); ch2.height=8.5; ch2.width=14; ch2.legend=None
try: ch2.series[0].graphicalProperties.solidFill="8E44AD"
except Exception: pass
ws4.add_chart(ch2,f"F{rt+3}")

ch3=PieChart(); ch3.title="Mix de Métodos de Pago"
d3=Reference(ws4,min_col=7,min_row=r1+1,max_row=r1+1+len(PAGOS)); cat3=Reference(ws4,min_col=6,min_row=r1+2,max_row=r1+1+len(PAGOS))
ch3.add_data(d3,titles_from_data=True); ch3.set_categories(cat3); ch3.height=8.5; ch3.width=11
ws4.add_chart(ch3,f"K{rt+3}")

leyenda(ws4,rt+21,14,"   ℹ️ Cambia el Mes (D6) y el Año (D7) — celdas amarillas — y todo el panel se recalcula.  Fuente de datos: hojas 1, 2 y 3.  Moneda: Soles (S/).")

# =========================================================
# Validación de datos + nombres definidos
# =========================================================
wb.defined_names.add(DefinedName("Lista_Categorias",attr_text=f"'0_Listas_y_Apoyo'!$A$4:$A${3+len(CATS)}"))
wb.defined_names.add(DefinedName("Lista_MetodosPago",attr_text=f"'0_Listas_y_Apoyo'!$B$4:$B${3+len(PAGOS)}"))
wb.defined_names.add(DefinedName("Lista_Proveedores",attr_text=f"'0_Listas_y_Apoyo'!$C$4:$C${3+len(PROVS)}"))
wb.defined_names.add(DefinedName("Lista_SKU",attr_text=f"{SH2}!$A${INI}:$A${S_FIN}"))

def dv(ws,formula,rng,titulo,msg):
    v=DataValidation(type="list",formula1=formula,allow_blank=True,showDropDown=False,
                     errorTitle="Valor no permitido",error=msg,promptTitle=titulo,prompt=msg,showErrorMessage=True,showInputMessage=True)
    ws.add_data_validation(v); v.add(rng)

dv(ws1,"=Lista_Proveedores",f"C{INI}:C{C_FIN}","Proveedor","Elige un proveedor de la lista maestra.")
dv(ws1,"=Lista_SKU",f"D{INI}:D{C_FIN}","Código/SKU","El SKU debe existir en 2_Control_de_Stock.")
dv(ws1,"=Lista_Categorias",f"F{INI}:F{C_FIN}","Categoría","Elige una categoría DAMABIS.")
dv(ws2,"=Lista_Categorias",f"C{INI}:C{S_FIN}","Categoría","Elige una categoría DAMABIS.")
dv(ws3,"=Lista_SKU",f"C{INI}:C{V_FIN}","Código/SKU","El SKU debe existir en 2_Control_de_Stock.")
dv(ws3,"=Lista_MetodosPago",f"H{INI}:H{V_FIN}","Método de Pago","Efectivo, Tarjeta, Transferencia, Yape/Plin o Crédito.")

vnum=DataValidation(type="whole",operator="greaterThan",formula1="0",allow_blank=True,showErrorMessage=True,
                    errorTitle="Cantidad inválida",error="Ingresa un número entero mayor a cero.")
ws3.add_data_validation(vnum); vnum.add(f"E{INI}:E{V_FIN}")
vnum2=DataValidation(type="whole",operator="greaterThan",formula1="0",allow_blank=True,showErrorMessage=True,
                     errorTitle="Cantidad inválida",error="Ingresa un número entero mayor a cero.")
ws1.add_data_validation(vnum2); vnum2.add(f"G{INI}:G{C_FIN}")
vmes=DataValidation(type="whole",operator="between",formula1="1",formula2="12",showErrorMessage=True,
                    errorTitle="Mes inválido",error="Ingresa un número de mes entre 1 y 12.")
ws4.add_data_validation(vmes); vmes.add("D6")

# Config de impresión
for ws in (ws1,ws2,ws3,ws4):
    ws.page_setup.orientation="landscape"; ws.page_setup.fitToWidth=1; ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.print_title_rows=f"{HDR}:{HDR}" if ws is not ws4 else None
wb.active=wb.sheetnames.index("4_Dashboard_DAMABIS")
ws0.sheet_state="visible"
wb.calculation.fullCalcOnLoad=True
wb.save(OUT)
print("OK ->",OUT)
